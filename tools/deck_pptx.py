"""발표 자료(deck)를 고칠 수 있는 PowerPoint(.pptx)로 내보내기.

    python tools/deck_pptx.py <발표> [--out 파일.pptx] [--font 글꼴이름]

빌드한 HTML(out/<발표>/wide/<발표>-wide.html)을 장마다 등장이 끝난 순간에 열어
글·도형·사진의 위치와 색을 그대로 읽어 와서 진짜 PowerPoint 개체로 옮겨요.

- 글은 글 상자(고칠 수 있음), 카드·막대·점·선은 도형, 사진은 그림으로 들어가요.
- 바탕(색·격자·그라디언트·창 틀·질감)은 고해상도 그림 한 장으로 깔아요.
- 강조색 도형(!!mo)과 바닥글은 모든 장에서 이름이 같아서 장을 넘길 때 PowerPoint '모핑' 전환으로 모양을 바꿔요.
- 글과 항목은 장이 열리면 저절로 차례대로 들어와요(나타내기 + 살짝 올라오기, 막대·선은 닦아내기). 클릭은 장마다 한 번.
- 발표 메모(notes)는 슬라이드 노트로 들어가요.
- 글꼴은 이 PC 에 있는 것 중에서 고르고(Pretendard → Noto Sans KR → 맑은 고딕, 명조는 Noto Serif KR → 바탕),
  Windows 에 PowerPoint 가 있으면 글꼴을 파일 안에 넣어(포함) 다른 PC 에서도 같은 글꼴로 보이게 해요.
"""
import asyncio
import io
import json
import os
import re
import sys
from pathlib import Path

from common import ROOT, load_json, out_dir

PX = 6350            # 1px (1920 기준) = 6350 EMU  → 슬라이드 13.333 × 7.5 in
PT = 0.5             # 1px = 0.5pt

# 장마다 읽어 오는 것 (브라우저 안에서 실행)
EXTRACT_JS = r"""
() => {
  const cam = document.querySelector('#cam'), C = cam.getBoundingClientRect(), K = C.width / 1920;
  const R = el => { const a = el.getBoundingClientRect(); return { x: (a.left - C.left) / K, y: (a.top - C.top) / K, w: a.width / K, h: a.height / K }; };
  const shown = el => { for (let e = el; e && e !== document.body; e = e.parentElement) { const s = getComputedStyle(e); if (s.visibility === 'hidden' || s.display === 'none' || +s.opacity < .04) return false; } const r = el.getBoundingClientRect(); return r.width > .5 && r.height > .5; };
  const alpha = el => { let a = 1; for (let e = el; e && e !== document.body; e = e.parentElement) a *= +getComputedStyle(e).opacity; return a; };
  const cur = [...document.querySelectorAll('.slide')].find(s => getComputedStyle(s).visibility !== 'hidden');
  const out = { texts: [], shapes: [], pics: [], mo: null, foot: [], bg: getComputedStyle(document.querySelector('#dk-bg')).backgroundColor };
  if (!cur) return out;
  out.index = +cur.dataset.i;
  const groupOf = el => { const g = el.closest('.cd, .cmp .side, .st, .sp, .bc, .bl li, .ag li, .cmp li, .qby, .hd, .bign'); return g ? g : el; };
  const gid = new Map(); let gn = 0;
  const G = el => { const g = groupOf(el); if (!gid.has(g)) gid.set(g, gn++); return gid.get(g); };
  // 글: 글자 노드를 따라가며 색·굵기·크기가 바뀌는 곳에서 끊음, 줄 상자(.mln)·<br>·\n 은 줄바꿈
  function runs(el) {
    const rs = []; let lastLine = null;
    const push = (t, s) => { const k = [s.color, s.fontWeight, s.fontSize, s.fontFamily, s.fontStyle].join('|'); const p = rs[rs.length - 1]; if (p && p.k === k) p.t += t; else rs.push({ k, t, color: s.color, weight: +s.fontWeight || 400, size: parseFloat(s.fontSize), family: s.fontFamily, italic: s.fontStyle === 'italic' }); };
    const walk = n => {
      if (n.nodeType === 3) { const pe = n.parentElement; if (!pe || pe.classList.contains('caret')) return; const ln = pe.closest('.mln'); if (ln && lastLine && ln !== lastLine) push('\n', getComputedStyle(pe)); if (ln) lastLine = ln; const s = getComputedStyle(pe); if (s.display === 'none') return; push(n.textContent, s); return; }
      if (n.nodeType !== 1) return;
      if (n.tagName === 'BR') { push('\n', getComputedStyle(n.parentElement)); return; }
      if (n.classList.contains('caret') || n.tagName === 'SVG' || n.tagName === 'svg') return;
      if (getComputedStyle(n).display === 'none') return;
      n.childNodes.forEach(walk);
    };
    walk(el);
    // pre-line: 줄 안 공백 정리
    rs.forEach(r => { r.t = r.t.replace(/[ \t]+/g, ' '); });
    return rs.filter(r => r.t.length);
  }
  const role = el => el.matches('.kick') ? 'kick' : el.matches('.htl') ? 'title' : el.matches('.big, .stm, .snum, .bign, .qtx, .qm') ? 'hero' : el.matches('.src, .credit') ? 'src' : 'body';
  const addText = el => {
    if (!shown(el)) return;
    const s = getComputedStyle(el), rr = runs(el); if (!rr.length || !rr.map(r => r.t).join('').trim()) return;
    const pr = R(el), pl = parseFloat(s.paddingLeft) || 0, pt = parseFloat(s.paddingTop) || 0, prr = parseFloat(s.paddingRight) || 0, pb = parseFloat(s.paddingBottom) || 0;
    const box = { x: pr.x + pl, y: pr.y + pt, w: Math.max(4, pr.w - pl - prr), h: Math.max(4, pr.h - pt - pb) };
    const lines = Math.round(box.h / (parseFloat(s.lineHeight) || parseFloat(s.fontSize) * 1.25));
    const center = s.display.includes('grid') || s.display.includes('flex') ? (s.justifyContent === 'center' || s.placeItems.includes('center')) : false;
    out.texts.push({ r: box, runs: rr, lines, center, align: center ? 'center' : s.textAlign, lh: parseFloat(s.lineHeight) || parseFloat(s.fontSize) * 1.25, ls: parseFloat(s.letterSpacing) || 0, size: parseFloat(s.fontSize),
      role: role(el), g: G(el), o: alpha(el), upper: s.textTransform === 'uppercase', nowrap: s.whiteSpace === 'nowrap' });
  };
  const TXT = '.kick, .htl, .big, .pg > .sub, .bd > .sub, .src, .stm, .snum, .bign, .nlab, .qm, .qtx, .qby b, .qby span, .ag li b, .ag li > span, .cd .no, .cd b, .cd > span, .st .v, .st .l, .chart .bv, .chart .bn, .cmp .tag, .cmp li, .steps .dt, .steps .sp > b, .steps .sp > span, .ptx, .bm b, .bm span, .logo span, .credit';
  cur.querySelectorAll(TXT).forEach(addText);
  cur.querySelectorAll('.bl li').forEach(li => { const b = li.querySelector(':scope > b'); if (b) li.querySelectorAll(':scope > b, :scope > span').forEach(addText); else addText(li); });
  // 도형
  const col = v => v && v !== 'rgba(0, 0, 0, 0)' && v !== 'transparent' ? v : null;
  const addBox = (el, kind) => {
    if (!shown(el)) return; const s = getComputedStyle(el), r = R(el);
    let ring = null; const m = /(rgba?\([^)]*\)|color\([^)]*\))\s+0px 0px 0px ([\d.]+)px inset/.exec(s.boxShadow) || /inset 0px 0px 0px ([\d.]+)px (rgba?\([^)]*\)|color\([^)]*\))/.exec(s.boxShadow);
    if (m) ring = m[1].startsWith('rgb') || m[1].startsWith('color') ? { c: m[1], w: +m[2] } : { c: m[2], w: +m[1] };
    const soft = s.boxShadow !== 'none' && s.boxShadow.split(/,(?![^(]*\))/).some(x => !/inset/.test(x) && (/(-?[\d.]+)px\s+(-?[\d.]+)px\s+([\d.]+)px/.exec(x.replace(/rgba?\([^)]*\)|color\([^)]*\)/g, '')) || [0, 0, 0, 0])[3] > 0);
    const fill = col(s.backgroundColor);
    if (!fill && !ring) return;
    out.shapes.push({ r, fill, ring, rad: parseFloat(s.borderTopLeftRadius) || 0, kind, g: G(el), o: alpha(el), soft });
  };
  cur.querySelectorAll('.cd, .cmp .side, .cmp .tag, .cmp .vs, .chart .bc > b, .chart .base, .steps .ln, .steps .ln > i, .steps .dt').forEach(el => addBox(el, el.matches('.chart .bc > b') ? 'bar' : el.matches('.chart .base, .steps .ln, .steps .ln > i') ? 'line' : el.matches('.steps .dt, .cmp .vs') ? 'dot' : 'box'));
  cur.querySelectorAll('.cmp .vs svg').forEach(sv => { if (!shown(sv)) return; out.shapes.push({ r: R(sv), fill: getComputedStyle(sv).fill, kind: 'arrow', rad: 0, g: G(sv), o: alpha(sv) }); });
  // 글머리 점(::before), 칸막이 선(border)
  cur.querySelectorAll('.bl li, .cmp li').forEach(li => { if (!shown(li)) return; const b = getComputedStyle(li, '::before'), lr = R(li); const w = parseFloat(b.width), h = parseFloat(b.height); if (!(w > 0 && h > 0)) return;
    out.shapes.push({ r: { x: lr.x + parseFloat(b.left || 0), y: lr.y + (parseFloat(b.top) || 0), w, h }, fill: col(b.backgroundColor), rad: parseFloat(b.borderTopLeftRadius) || 0, kind: 'bullet', g: G(li), o: alpha(li) * (+b.opacity || 1) }); });
  cur.querySelectorAll('.st').forEach(st => { if (!shown(st)) return; const s = getComputedStyle(st); if (parseFloat(s.borderLeftWidth) > 0) { const r = R(st); out.shapes.push({ r: { x: r.x, y: r.y, w: parseFloat(s.borderLeftWidth), h: r.h }, fill: s.borderLeftColor, rad: 0, kind: 'line', g: G(st), o: alpha(st) }); } });
  cur.querySelectorAll('.ag li').forEach(li => { if (!shown(li)) return; const s = getComputedStyle(li); if (parseFloat(s.borderBottomWidth) > 0) { const r = R(li); out.shapes.push({ r: { x: r.x, y: r.y + r.h - parseFloat(s.borderBottomWidth), w: r.w, h: parseFloat(s.borderBottomWidth) }, fill: s.borderBottomColor, rad: 0, kind: 'line', g: G(li), o: alpha(li) }); } });
  // 사진·로고 (그림으로 찍어 넣음)
  cur.querySelectorAll('.ph, .bm img, .logo img, .qby img').forEach(el => { if (shown(el)) out.pics.push({ r: R(el), g: G(el), kind: el.matches('.ph') ? 'photo' : 'logo' }); });
  // 따라다니는 강조 도형
  const mo = document.querySelector('#dk-mo'), ms = getComputedStyle(mo);
  if (shown(mo) && +ms.opacity > .05) out.mo = { r: R(mo), fill: ms.backgroundColor, rad: parseFloat(ms.borderTopLeftRadius) || 0 };
  // 바닥글
  const ft = document.querySelector('#dk-ft');
  if (ft && shown(ft)) ft.querySelectorAll('.brand span, .ttl, .no').forEach((el, k) => { if (!shown(el)) return; const s = getComputedStyle(el), rr = runs(el); if (rr.length) out.foot.push({ nowrap: true, lines: 1, name: ['brand', 'ttl', 'no'][el.matches('.brand span') ? 0 : el.matches('.ttl') ? 1 : 2], r: R(el), runs: rr, align: s.textAlign, lh: parseFloat(s.lineHeight) || 20, ls: parseFloat(s.letterSpacing) || 0, size: parseFloat(s.fontSize) }); });
  return out;
}
"""


# ------------------------------------------------------------------ 색
def parse_color(v):
    """CSS 색 → (RRGGBB, 투명도 0~1). rgb()/rgba()/color(srgb …)"""
    if not v:
        return None, 1
    v = v.strip()
    m = re.match(r"rgba?\(\s*([\d.]+)[ ,]+([\d.]+)[ ,]+([\d.]+)(?:\s*[,/]\s*([\d.]+%?))?\s*\)", v)
    if m:
        r, g, b = (int(round(float(x))) for x in m.groups()[:3])
        a = m.group(4)
        a = (float(a[:-1]) / 100 if a and a.endswith("%") else float(a)) if a else 1
        return f"{r:02X}{g:02X}{b:02X}", a
    m = re.match(r"color\(srgb\s+([\d.e-]+)\s+([\d.e-]+)\s+([\d.e-]+)(?:\s*/\s*([\d.]+%?))?\)", v)
    if m:
        r, g, b = (max(0, min(255, int(round(float(x) * 255)))) for x in m.groups()[:3])
        a = m.group(4)
        a = (float(a[:-1]) / 100 if a and a.endswith("%") else float(a)) if a else 1
        return f"{r:02X}{g:02X}{b:02X}", a
    m = re.match(r"#([0-9a-fA-F]{6})", v)
    if m:
        return m.group(1).upper(), 1
    return None, 1


def blend(hex_fg, a, hex_bg):
    """반투명 색을 바탕색에 섞은 불투명 색 (글자는 투명도를 못 써서)"""
    if a >= .999 or not hex_bg:
        return hex_fg
    f = [int(hex_fg[i:i + 2], 16) for i in (0, 2, 4)]
    b = [int(hex_bg[i:i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(f[i] * a + b[i] * (1 - a)):02X}" for i in range(3))


# ------------------------------------------------------------------ 글꼴
def installed_fonts():
    names = set()
    if os.name == "nt":
        try:
            import winreg
            for root in (winreg.HKEY_LOCAL_MACHINE, winreg.HKEY_CURRENT_USER):
                try:
                    k = winreg.OpenKey(root, r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts")
                except OSError:
                    continue
                i = 0
                while True:
                    try:
                        n = winreg.EnumValue(k, i)[0]
                    except OSError:
                        break
                    names.add(re.sub(r"\s*\((TrueType|OpenType)\)$", "", n))
                    i += 1
        except Exception:  # noqa: BLE001
            pass
    return names


class Fonts:
    """CSS 글꼴(이름·굵기) → PowerPoint 글꼴 이름과 굵게 여부"""
    SANS_W = {900: "Black", 800: "ExtraBold", 700: "Bold", 600: "SemiBold", 500: "Medium", 400: "Regular", 300: "Light"}

    def __init__(self, force=None, files=None):
        self.have = installed_fonts()
        self.force = force
        self.used = set()
        self.files = {k.lower(): v.lower() for k, v in (files or {}).items()}   # 브랜드·스타일 글꼴 이름 → 파일 이름

    def kind(self, family):
        """첫 번째 실제 글꼴 이름으로 성격을 정함 (serif·sans-serif 같은 대체 이름은 건너뜀)"""
        for tok in (family or "").split(","):
            n = tok.strip().strip('"\'').lower()
            if not n or n in ("serif", "sans-serif", "monospace", "cursive", "system-ui"):
                continue
            key = self.files.get(n, "") + " " + n
            if any(k in key for k in ("caveat", "nanumpen", "nanum pen", "nanumbrush", "nanum brush", "gaegu")):
                return "hand"
            if any(k in key for k in ("mono", "galmuri", "silkscreen")):
                return "mono"
            if ("serif" in key and "sans" not in key) or any(k in key for k in ("batang", "myeongjo", "cinzel")):
                return "serif"
            return "sans"
        return "sans"

    def has(self, n):
        return not self.have or n in self.have   # 목록을 못 읽으면(리눅스 등) 있다고 봄

    def pick(self, family, weight):
        k = self.kind(family)
        serif, hand, mono = k == "serif", k == "hand", k == "mono"
        w = min(self.SANS_W, key=lambda x: abs(x - (weight or 400)))
        if self.force:
            name, bold = self.force, weight >= 600
        elif serif:
            name, bold = ("Noto Serif KR", weight >= 600) if self.has("Noto Serif KR") else ("바탕", weight >= 600)
        elif hand and self.has("Nanum Pen"):
            name, bold = "Nanum Pen", False
        elif mono and self.has("D2Coding"):
            name, bold = "D2Coding", weight >= 600
        elif self.has(f"Pretendard {self.SANS_W[w]}") or self.has("Pretendard Black"):
            name = "Pretendard" if w == 400 else f"Pretendard {self.SANS_W[w]}"
            bold = False
        elif self.has("Noto Sans KR"):
            name, bold = "Noto Sans KR", weight >= 600
        else:
            name, bold = "맑은 고딕", weight >= 600
        self.used.add(name)
        return name, bold


# ------------------------------------------------------------------ 브라우저에서 읽기
async def capture(name, work):
    from playwright.async_api import async_playwright
    from PIL import Image
    html = out_dir(name, "wide") / f"{name}-wide.html"
    if not html.exists():
        sys.exit(f"{html.relative_to(ROOT)} 이 없어요. 먼저 python tools/build.py {name} 을 실행하세요.")
    async with async_playwright() as p:
        browser = await p.chromium.launch()
        page = await browser.new_page(viewport={"width": 1920, "height": 1080}, device_scale_factor=2)
        await page.add_init_script("window.__RENDER=true;")
        await page.goto(html.as_uri())
        await page.wait_for_function("window.__ready===true", timeout=60000)
        stops = await page.evaluate("window.__STOPS || []")
        if not stops:
            await browser.close()
            sys.exit("발표 자료(deck 템플릿)가 아니에요.")
        # 바탕: 장·도형·바닥글을 숨기고 찍음
        await page.evaluate(f"window.__seek({stops[0]})")
        await page.add_style_tag(content=".slide,#dk-mo,#dk-ft{visibility:hidden!important} .ms-hide-pg .pg, .ms-hide-pg #dk-ft{visibility:hidden!important}")
        bg = work / "bg.jpg"
        Image.open(io.BytesIO(await page.screenshot(type="png"))).convert("RGB").save(bg, quality=92)
        await page.evaluate("document.querySelectorAll('style').forEach(s => { if (s.textContent.startsWith('.slide,#dk-mo')) s.textContent = '.ms-hide-pg .pg, .ms-hide-pg #dk-ft{visibility:hidden!important}'; })")
        slides = []
        for i, t in enumerate(stops):
            await page.evaluate(f"window.__seek({t})")
            await page.wait_for_timeout(40)
            d = await page.evaluate(EXTRACT_JS)
            # 사진은 글을 숨기고 그 자리만 찍음
            if d["pics"]:
                await page.evaluate("document.body.classList.add('ms-hide-pg')")
                for k, pc in enumerate(d["pics"]):
                    r = pc["r"]
                    clip = {"x": max(0, r["x"]), "y": max(0, r["y"]), "width": min(1920, r["x"] + r["w"]) - max(0, r["x"]), "height": min(1080, r["y"] + r["h"]) - max(0, r["y"])}
                    if clip["width"] < 2 or clip["height"] < 2:
                        continue
                    if pc["kind"] == "logo":   # 로고는 글 숨김과 상관없이 보여야 해서 따로
                        await page.evaluate("document.body.classList.remove('ms-hide-pg')")
                    f = work / f"s{i}_p{k}.png"
                    await page.screenshot(path=str(f), clip=clip)
                    pc["file"], pc["r"] = str(f), {"x": clip["x"], "y": clip["y"], "w": clip["width"], "h": clip["height"]}
                    if pc["kind"] == "logo":
                        await page.evaluate("document.body.classList.add('ms-hide-pg')")
                await page.evaluate("document.body.classList.remove('ms-hide-pg')")
            slides.append(d)
        await browser.close()
    return bg, slides


# ------------------------------------------------------------------ PowerPoint 만들기
NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def emu(px):
    from pptx.util import Emu
    return Emu(int(round(px * PX)))


def add_text(slide, t, fonts, bg_hex, name=None):
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    r = t["r"]
    align = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "end": PP_ALIGN.RIGHT}.get(t["align"], PP_ALIGN.LEFT)
    # PowerPoint 글꼴 폭이 조금 달라도 줄이 더 꺾이지 않게 상자를 넓힘 (정렬 방향 반대쪽으로)
    extra = max(r["w"] * .08, t["size"] * 1.2)
    x, w = r["x"], r["w"] + extra
    if align == PP_ALIGN.CENTER:
        x -= extra / 2
    elif align == PP_ALIGN.RIGHT:
        x -= extra
    single = t.get("nowrap") or (t.get("lines", 2) <= 1 and "\n" not in "".join(x["t"] for x in t["runs"]).strip())
    if single or t.get("center"):
        # 한 줄 글: PowerPoint 의 '정확히' 줄 간격은 글자를 위로 밀어서, 기본 간격 + 가운데 맞춤으로 원래 줄 가운데에 둠
        cy = r["y"] + (r["h"] if t.get("center") else min(r["h"], t["lh"])) / 2
        hh = max(r["h"] if t.get("center") else 0, t["size"] * 1.5)
        box = slide.shapes.add_textbox(emu(x), emu(cy - hh / 2), emu(w), emu(hh))
    else:
        box = slide.shapes.add_textbox(emu(x), emu(r["y"]), emu(w), emu(max(r["h"], t["lh"])))
    if name:
        box.name = name
    tf = box.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = not single   # 화면에서 한 줄이던 글은 PowerPoint 에서도 꺾이지 않게
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE if (single or t.get("center")) else MSO_ANCHOR.TOP
    tf.auto_size = None
    # 줄바꿈으로 문단을 나눔
    paras = [[]]
    for run in t["runs"]:
        parts = run["t"].split("\n")
        for k, part in enumerate(parts):
            if k:
                paras.append([])
            if part:
                paras[-1].append(dict(run, t=part))
    paras = [p for i, p in enumerate(paras) if p or 0 < i < len(paras) - 1]
    # 앞뒤 공백 정리
    for p in paras:
        if p:
            p[0]["t"] = p[0]["t"].lstrip()
            p[-1]["t"] = p[-1]["t"].rstrip()
    for i, pr in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        if not (single or t.get("center")):
            para.line_spacing = Pt(t["lh"] * PT)
        para.space_before = para.space_after = Pt(0)
        for run in pr:
            if not run["t"]:
                continue
            rr = para.add_run()
            rr.text = run["t"].upper() if t.get("upper") else run["t"]
            f = rr.font
            fname, bold = fonts.pick(run["family"], run["weight"])
            f.name = fname
            f.bold = bold
            f.italic = run.get("italic") or None
            f.size = Pt(round(run["size"] * PT * 2) / 2)
            c, a = parse_color(run["color"])
            if c:
                f.color.rgb = RGBColor.from_string(blend(c, a * t.get("o", 1), bg_hex))
            # 한글 글꼴도 같은 이름으로 (동아시아 글꼴 칸)
            rPr = rr._r.get_or_add_rPr()
            for tag in ("a:ea", "a:cs"):
                el = rPr.find(tag.replace("a:", "{%s}" % NS["a"]))
                if el is None:
                    from lxml import etree
                    el = etree.SubElement(rPr, "{%s}%s" % (NS["a"], tag[2:]))
                el.set("typeface", fname)
            if t.get("ls"):
                rPr.set("spc", str(int(round(t["ls"] * PT * 100))))
    return box


def add_shape(slide, s, bg_hex, name=None):
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    r = s["r"]
    kind = s.get("kind")
    rad = s.get("rad") or 0
    if kind == "arrow":
        shp = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, emu(r["x"] + r["w"] * .12), emu(r["y"] + r["h"] * .28), emu(r["w"] * .76), emu(r["h"] * .44))
    elif rad >= min(r["w"], r["h"]) / 2 - .5 and abs(r["w"] - r["h"]) < 1.5:
        shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, emu(r["x"]), emu(r["y"]), emu(r["w"]), emu(r["h"]))
    elif rad > .5:
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, emu(r["x"]), emu(r["y"]), emu(r["w"]), emu(r["h"]))
        shp.adjustments[0] = max(0, min(.5, rad / max(1, min(r["w"], r["h"]))))
    else:
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, emu(r["x"]), emu(r["y"]), emu(r["w"]), emu(r["h"]))
    if name:
        shp.name = name
    c, a = parse_color(s.get("fill"))
    a *= s.get("o", 1)
    if c:
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(c)
        if a < .999:
            set_alpha(shp.fill._xPr.find("{%s}solidFill" % NS["a"])[0], a)
    else:
        shp.fill.background()
    ring = s.get("ring")
    if ring:
        rc, ra = parse_color(ring["c"])
        shp.line.color.rgb = RGBColor.from_string(rc)
        shp.line.width = Pt(ring["w"] * PT)
        if ra < .999:
            set_alpha(shp.line._get_or_add_ln().find("{%s}solidFill" % NS["a"])[0], ra)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    if s.get("soft"):
        add_soft_shadow(shp)
    if shp.has_text_frame:
        shp.text_frame.text = ""
    return shp


def set_alpha(clr, a):
    from lxml import etree
    for old in clr.findall("{%s}alpha" % NS["a"]):
        clr.remove(old)
    el = etree.SubElement(clr, "{%s}alpha" % NS["a"])
    el.set("val", str(int(round(max(0, min(1, a)) * 100000))))


def add_soft_shadow(shp):
    from lxml import etree
    spPr = shp._element.spPr
    for old in spPr.findall("{%s}effectLst" % NS["a"]):
        spPr.remove(old)
    eff = etree.SubElement(spPr, "{%s}effectLst" % NS["a"])
    sh = etree.SubElement(eff, "{%s}outerShdw" % NS["a"], blurRad="190500", dist="38100", dir="5400000", algn="t", rotWithShape="0")
    c = etree.SubElement(sh, "{%s}srgbClr" % NS["a"], val="000000")
    etree.SubElement(c, "{%s}alpha" % NS["a"], val="16000")


def anim_xml(items, ids):
    """장이 열리면 저절로 차례대로 들어오는 애니메이션. items = [(spid, 시작초, 종류)]"""
    if not items:
        return ""
    n = [3]

    def nid():
        n[0] += 1
        return n[0]
    pars = []
    for spid, start, kind in items:
        d = 500 if kind != "wipe" else 600
        a, b, c = nid(), nid(), nid()
        tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
        setv = (f'<p:set><p:cBhvr><p:cTn id="{b}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}'
                f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>')
        if kind == "wipe-up" or kind == "wipe-left":
            filt = "wipe(up)" if kind == "wipe-up" else "wipe(right)"
            body = setv + f'<p:animEffect transition="in" filter="{filt}"><p:cBhvr><p:cTn id="{c}" dur="{d}"/>{tgt}</p:cBhvr></p:animEffect>'
            preset, sub = 22, (4 if kind == "wipe-up" else 8)
        else:
            e = nid()
            body = (setv + f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{c}" dur="{d}"/>{tgt}</p:cBhvr></p:animEffect>'
                    f'<p:anim calcmode="lin" valueType="num"><p:cBhvr additive="base"><p:cTn id="{e}" dur="{d}" decel="100000" fill="hold"/>{tgt}'
                    f'<p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst>'
                    f'<p:tav tm="0"><p:val><p:strVal val="#ppt_y+0.025"/></p:val></p:tav><p:tav tm="100000"><p:val><p:strVal val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim>')
            preset, sub = 42, 0   # 올라오며 나타내기(Ascend 계열)
        pars.append(f'<p:par><p:cTn id="{a}" presetID="{preset}" presetClass="entr" presetSubtype="{sub}" fill="hold" grpId="0" nodeType="withEffect">'
                    f'<p:stCondLst><p:cond delay="{int(start * 1000)}"/></p:stCondLst><p:childTnLst>{body}</p:childTnLst></p:cTn></p:par>')
    bld = "".join(f'<p:bldP spid="{s}" grpId="0"/>' if t != "pic" else "" for s, t in ids)
    bldg = "".join(f'<p:bldGraphic spid="{s}" grpId="0"><p:bldAsOne/></p:bldGraphic>' for s, t in ids if t == "pic")
    return (f'<p:timing xmlns:p="{NS["p"]}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
            f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
            f'<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="{nid()}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            + "".join(pars) +
            '</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn>'
            '<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
            '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
            f'</p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst>{bld}{bldg}</p:bldLst></p:timing>')


MORPH = ('<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
         '<mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" Requires="p159">'
         '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" spd="slow" p14:dur="800"><p159:morph option="byObject"/></p:transition></mc:Choice>'
         '<mc:Fallback><p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="med"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>')


def build(name, bg, slides, out, font=None):
    from lxml import etree
    from pptx import Presentation
    from pptx.util import Emu
    v = load_json(ROOT / "videos" / name / "video.json")
    content = v.get("content") or {}
    notes = [str((s or {}).get("notes") or "").strip() for s in content.get("slides") or []]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(1920 * PX), Emu(1080 * PX)
    blank = prs.slide_layouts[6]
    files = {}
    b = load_json(ROOT / "brands" / v["brand"] / "brand.json")
    for f in b.get("fonts") or []:
        files[f["family"]] = Path(f["file"]).name
    if v.get("style") and (ROOT / "styles" / v["style"] / "style.json").exists():
        for f in load_json(ROOT / "styles" / v["style"] / "style.json").get("fonts") or []:
            files[f["family"]] = Path(f["file"]).name
    fonts = Fonts(font, files)
    prs.core_properties.title = v.get("title") or name
    for i, d in enumerate(slides):
        sl = prs.slides.add_slide(blank)
        bg_hex, _ = parse_color(d.get("bg"))
        # 바탕 그림 (모든 장이 같은 이름 → 모핑 때 가만히 있음)
        pic = sl.shapes.add_picture(str(bg), 0, 0, prs.slide_width, prs.slide_height)
        pic.name = "!!bg"
        if d.get("mo"):
            add_shape(sl, {"r": d["mo"]["r"], "fill": d["mo"]["fill"], "rad": d["mo"]["rad"], "kind": "box"}, bg_hex, name="!!mo")
        # 구분 장처럼 도형이 화면을 덮으면 글 위 바탕색은 그 도형 색
        mo_full = d.get("mo") and d["mo"]["r"]["w"] > 1800 and d["mo"]["r"]["h"] > 1000
        txt_bg = parse_color(d["mo"]["fill"])[0] if mo_full else bg_hex
        anim, ids = [], []
        groups = {}
        order = []

        def note(gid):
            if gid not in groups:
                groups[gid] = len(order)
                order.append(gid)
        for t in d["texts"]:
            note(("t", t["role"]) if t["role"] in ("kick", "title") else t["g"])
        for s in d["shapes"]:
            note(s["g"])
        for p in d["pics"]:
            note(p["g"])
        # 시작 시각: 꼬리표 0, 제목 0.1, 그 다음 묶음마다 0.15초씩
        def start_of(key, role=None):
            if role == "kick":
                return 0
            if role == "title":
                return .1
            k = [g for g in order if not (isinstance(g, tuple))].index(key) if key in order and not isinstance(key, tuple) else 0
            return .45 + min(k, 10) * .15
        for p in d["pics"]:
            if not p.get("file"):
                continue
            pp = sl.shapes.add_picture(p["file"], emu(p["r"]["x"]), emu(p["r"]["y"]), emu(p["r"]["w"]), emu(p["r"]["h"]))
            if p["kind"] == "photo":
                pp.name = f"사진 {len(ids) + 1}"
                anim.append((pp.shape_id, 0, "fade"))
            ids.append((pp.shape_id, "pic"))
        for s in d["shapes"]:
            shp = add_shape(sl, s, txt_bg)
            kind = "wipe-up" if s["kind"] == "bar" else "wipe-left" if s["kind"] == "line" and s["r"]["w"] > s["r"]["h"] * 3 else "fade"
            anim.append((shp.shape_id, start_of(s["g"]) + (.1 if s["kind"] == "bar" else 0), kind))
            ids.append((shp.shape_id, "sp"))
        for t in d["texts"]:
            box = add_text(sl, t, fonts, txt_bg)
            st = start_of(t["g"], t["role"]) if t["role"] in ("kick", "title") else start_of(t["g"])
            if t["role"] == "hero":
                st = .15
            if t["role"] == "src":
                st = max(st, .9)
            anim.append((box.shape_id, st, "fade"))
            ids.append((box.shape_id, "sp"))
        for k, f in enumerate(d["foot"]):
            add_text(sl, dict(f, role="foot", o=1), fonts, bg_hex, name=f"!!ft-{f['name']}")
        # 장이 열리면 저절로 (모핑이 끝난 뒤)
        anim.sort(key=lambda x: x[1])
        if anim:
            sl._element.append(etree.fromstring(anim_xml(anim, ids)))
        if i:
            tr = etree.fromstring(MORPH)
            timing = sl._element.find("{%s}timing" % NS["p"])
            if timing is not None:
                timing.addprevious(tr)
            else:
                sl._element.append(tr)
        idx = d.get("index", i)
        if idx < len(notes) and notes[idx]:
            sl.notes_slide.notes_text_frame.text = notes[idx]
    prs.save(out)
    return sorted(fonts.used)


def embed_fonts(path):
    """Windows PowerPoint 로 열어 글꼴을 파일에 넣어(포함) 다시 저장. 안 되면 그대로 둠"""
    if os.name != "nt":
        return False
    try:
        import win32com.client  # noqa: F401
    except ImportError:
        return _embed_ps(path)
    try:
        import pythoncom
        import win32com.client
        pythoncom.CoInitialize()
        app = win32com.client.DispatchEx("PowerPoint.Application")
        pres = app.Presentations.Open(str(path), True, False, False)
        tmp = str(Path(path).with_suffix(".tmp.pptx"))
        pres.SaveAs(tmp, 24, True)   # ppSaveAsOpenXMLPresentation, EmbedTrueTypeFonts
        pres.Close()
        app.Quit()
        os.replace(tmp, path)
        return True
    except Exception:  # noqa: BLE001
        return _embed_ps(path)


def _embed_ps(path):
    import subprocess
    ps = ("$ErrorActionPreference='Stop'; $a=New-Object -ComObject PowerPoint.Application; "
          f"$p=$a.Presentations.Open('{path}', -1, 0, 0); $t='{Path(path).with_suffix('.tmp.pptx')}'; "
          "$p.SaveAs($t, 24, -1); $p.Close(); $a.Quit(); "
          f"Move-Item -Force $t '{path}'")
    try:
        r = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, timeout=180)
        return r.returncode == 0
    except Exception:  # noqa: BLE001
        return False


def export(name, out=None, font=None, embed=True):
    import tempfile
    work = Path(tempfile.mkdtemp(prefix="ms-pptx-"))
    try:
        bg, slides = asyncio.run(capture(name, work))
        out = Path(out) if out else out_dir(name, "wide") / f"{name}.pptx"
        used = build(name, bg, slides, out, font)
        embedded = embed_fonts(out.resolve()) if embed else False
        print(f"PPT {len(slides)}장 → {out.relative_to(ROOT) if out.is_relative_to(ROOT) else out}")
        print(f"  글꼴: {', '.join(used)}" + (" (파일에 포함)" if embedded else ""))
        return out, len(slides), used, embedded
    finally:
        import shutil
        shutil.rmtree(work, ignore_errors=True)


def main():
    a = sys.argv[1:]
    if not a or a[0].startswith("-"):
        sys.exit(__doc__)
    out = a[a.index("--out") + 1] if "--out" in a else None
    font = a[a.index("--font") + 1] if "--font" in a else None
    export(a[0], out, font, embed="--no-embed" not in a)


if __name__ == "__main__":
    main()
